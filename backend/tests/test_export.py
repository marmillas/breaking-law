import os
import subprocess
import uuid
import pytest
from unittest.mock import MagicMock, patch
from docx import Document

from breaking_law.documents.export import ExportService


def test_export_worker_actor_importable():
    # Importing the module should register the actor without errors
    from breaking_law.documents.worker_export import export_document_version
    assert export_document_version is not None
    assert hasattr(export_document_version, "send")


@pytest.fixture
def mock_storage():
    storage = MagicMock()
    storage.upload_fileobj.return_value = True
    storage.generate_presigned_url.return_value = "https://example.com/signed-url"
    return storage


def test_generate_docx_from_structured_content(tmp_path):
    draft = {
        "title": "Test Document",
        "sections": [
            {"heading": "Section 1", "paragraphs": ["First paragraph.", "Second paragraph."]},
            {"heading": "Section 2", "paragraphs": ["Another paragraph."]},
        ],
    }
    output_path = tmp_path / "output.docx"
    service = ExportService(storage=MagicMock())
    service.generate_docx(draft, str(output_path))

    assert output_path.exists()
    doc = Document(str(output_path))
    paragraphs = [p.text for p in doc.paragraphs]
    assert "Test Document" in paragraphs
    assert "Section 1" in paragraphs
    assert "First paragraph." in paragraphs
    assert "Section 2" in paragraphs
    assert "Another paragraph." in paragraphs


def test_generate_docx_with_branding(tmp_path):
    draft = {"title": "Branded Doc", "text": "Content here."}
    output_path = tmp_path / "branded.docx"
    service = ExportService(storage=MagicMock())
    branding = {"law_firm_name": "Acme Legal"}
    service.generate_docx(draft, str(output_path), branding=branding)

    doc = Document(str(output_path))
    header_text = doc.sections[0].header.paragraphs[0].text
    assert header_text == "Acme Legal"


def test_generate_docx_with_template(tmp_path):
    template_path = tmp_path / "template.docx"
    template_doc = Document()
    template_doc.add_paragraph("Template Header")
    template_doc.save(str(template_path))

    draft = {"title": "From Template", "text": "Body text."}
    output_path = tmp_path / "from_template.docx"
    service = ExportService(storage=MagicMock())
    service.generate_docx(draft, str(output_path), template_path=str(template_path))

    doc = Document(str(output_path))
    paragraphs = [p.text for p in doc.paragraphs]
    assert "Template Header" in paragraphs
    assert "From Template" in paragraphs
    assert "Body text." in paragraphs


def test_convert_to_pdf_success(tmp_path):
    docx_path = tmp_path / "input.docx"
    Document().save(str(docx_path))

    expected_pdf = tmp_path / "input.pdf"

    def fake_subprocess(*args, **kwargs):
        # Simulate LibreOffice creating the PDF
        expected_pdf.write_text("fake pdf content")
        return MagicMock(returncode=0)

    service = ExportService(storage=MagicMock())
    with patch("subprocess.run", side_effect=fake_subprocess):
        pdf_path = service.convert_to_pdf(str(docx_path), str(tmp_path))

    assert pdf_path == str(expected_pdf)
    assert os.path.exists(pdf_path)


def test_convert_to_pdf_libreoffice_not_found():
    service = ExportService(storage=MagicMock(), libreoffice_path="/nonexistent/libreoffice")
    with patch("subprocess.run", side_effect=FileNotFoundError("No such file")):
        pdf_path = service.convert_to_pdf("/fake/input.docx", "/fake/output")
    assert pdf_path is None


def test_convert_to_pdf_called_process_error():
    service = ExportService(storage=MagicMock())
    with patch("subprocess.run", side_effect=subprocess.CalledProcessError(1, "libreoffice")):
        pdf_path = service.convert_to_pdf("/fake/input.docx", "/fake/output")
    assert pdf_path is None


@pytest.mark.asyncio
async def test_export_document_returns_signed_urls(mock_storage):
    service = ExportService(storage=mock_storage)
    draft = {"title": "Export Me", "text": "Some content."}
    version_id = uuid.uuid4()
    law_firm_id = uuid.uuid4()

    # Mock PDF conversion success by creating the PDF inside the temp dir
    original_convert = service.convert_to_pdf

    def mock_convert(docx_path, output_dir):
        base = os.path.splitext(os.path.basename(docx_path))[0] + ".pdf"
        pdf_path = os.path.join(output_dir, base)
        with open(pdf_path, "w") as f:
            f.write("pdf")
        return pdf_path

    service.convert_to_pdf = mock_convert

    result = await service.export_document(draft, version_id, law_firm_id)

    assert result["docx_url"] == "https://example.com/signed-url"
    assert result["pdf_url"] == "https://example.com/signed-url"
    assert result["docx_storage_key"] is not None
    assert result["pdf_storage_key"] is not None
    assert result["warning"] is None


@pytest.mark.asyncio
async def test_export_document_pdf_fallback_when_libreoffice_missing(mock_storage):
    service = ExportService(storage=mock_storage, libreoffice_path="/nonexistent/libreoffice")
    draft = {"title": "No PDF", "text": "Content."}
    version_id = uuid.uuid4()
    law_firm_id = uuid.uuid4()

    with patch("subprocess.run", side_effect=FileNotFoundError("No such file")):
        result = await service.export_document(draft, version_id, law_firm_id)

    assert result["docx_url"] == "https://example.com/signed-url"
    assert result["pdf_url"] is None
    assert result["pdf_storage_key"] is None
    assert result["warning"] is not None
    assert "LibreOffice" in result["warning"]
