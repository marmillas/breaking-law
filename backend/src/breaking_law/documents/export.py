"""
Export infrastructure for legal platform.

Provides document generation (DOCX) and PDF conversion using LibreOffice,
with upload to object storage and signed URL generation.
"""

import os
import uuid
import tempfile
import subprocess
from typing import Optional, Dict, Any

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from breaking_law.documents.storage import Storage


class ExportService:
    """Service for exporting document drafts to DOCX and PDF formats."""

    def __init__(
        self,
        storage: Storage,
        libreoffice_path: str = "/usr/bin/libreoffice",
        template_dir: Optional[str] = None,
    ):
        """
        Initialize the export service.

        Args:
            storage: S3-compatible storage adapter for artifact upload.
            libreoffice_path: Path to the LibreOffice executable for PDF conversion.
            template_dir: Optional directory containing DOCX templates.
        """
        self.storage = storage
        self.libreoffice_path = libreoffice_path
        self.template_dir = template_dir

    def _get_template_path(self, template_name: Optional[str] = None) -> Optional[str]:
        """Resolve a template name to an absolute file path."""
        if not template_name or not self.template_dir:
            return None
        path = os.path.join(self.template_dir, template_name)
        if os.path.isfile(path):
            return path
        return None

    def generate_docx(
        self,
        draft: Dict[str, Any],
        output_path: str,
        template_path: Optional[str] = None,
        branding: Optional[Dict[str, Any]] = None,
    ) -> str:
        """
        Generate a DOCX file from a structured draft.

        Args:
            draft: Dictionary containing document content.
                Expected keys:
                - title (str): Document title.
                - sections (list): Each section may have 'heading' (str)
                  and 'paragraphs' (list of str).
                - text (str): Fallback plain text if sections are absent.
            output_path: File path to write the DOCX to.
            template_path: Optional path to a DOCX template file.
            branding: Optional dict with law firm letterhead details,
                e.g. {"law_firm_name": "Acme Legal"}.

        Returns:
            The output path.
        """
        if template_path and os.path.isfile(template_path):
            doc = Document(template_path)
        else:
            doc = Document()

        # Apply branding / letterhead if provided
        if branding:
            law_firm_name = branding.get("law_firm_name", "")
            if law_firm_name:
                header = doc.sections[0].header
                if header.paragraphs:
                    header.paragraphs[0].text = law_firm_name
                else:
                    para = header.add_paragraph()
                    para.text = law_firm_name
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Title
        title = draft.get("title", "Untitled Document")
        doc.add_heading(title, level=0)

        # Sections
        sections = draft.get("sections", [])
        if sections:
            for section in sections:
                heading = section.get("heading")
                if heading:
                    doc.add_heading(heading, level=1)
                for para_text in section.get("paragraphs", []):
                    if para_text.strip():
                        doc.add_paragraph(para_text)
        else:
            # Fallback to plain text splitting on newlines
            text = draft.get("text", "")
            for line in text.split("\n"):
                if line.strip():
                    doc.add_paragraph(line)

        doc.save(output_path)
        return output_path

    def convert_to_pdf(self, docx_path: str, output_dir: str) -> Optional[str]:
        """
        Convert a DOCX file to PDF using LibreOffice CLI.

        Args:
            docx_path: Path to the source DOCX file.
            output_dir: Directory to write the PDF into.

        Returns:
            Path to the generated PDF, or None if conversion failed.
        """
        try:
            cmd = [
                self.libreoffice_path,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                output_dir,
                docx_path,
            ]
            subprocess.run(cmd, check=True, capture_output=True, timeout=120)
            base_name = os.path.splitext(os.path.basename(docx_path))[0] + ".pdf"
            pdf_path = os.path.join(output_dir, base_name)
            if os.path.exists(pdf_path):
                return pdf_path
            return None
        except (subprocess.CalledProcessError, FileNotFoundError, OSError):
            return None

    async def export_document(
        self,
        draft: Dict[str, Any],
        version_id: uuid.UUID,
        law_firm_id: uuid.UUID,
        template_name: Optional[str] = None,
        branding: Optional[Dict[str, Any]] = None,
    ) -> Dict[str, Any]:
        """
        Export a draft to DOCX and PDF, upload both, and return signed URLs.

        Args:
            draft: Structured document content.
            version_id: Document version UUID used for storage path.
            law_firm_id: Tenant UUID used for storage path.
            template_name: Optional template file name (resolved against template_dir).
            branding: Optional letterhead/branding details.

        Returns:
            Dictionary with:
                - docx_url (Optional[str])
                - pdf_url (Optional[str])
                - docx_storage_key (str)
                - pdf_storage_key (Optional[str])
                - warning (Optional[str])
        """
        template_path = self._get_template_path(template_name)

        with tempfile.TemporaryDirectory() as tmpdir:
            docx_filename = f"{version_id}.docx"
            docx_path = os.path.join(tmpdir, docx_filename)
            self.generate_docx(draft, docx_path, template_path=template_path, branding=branding)

            # Upload DOCX
            docx_key = f"{law_firm_id}/exports/{version_id}/{docx_filename}"
            with open(docx_path, "rb") as f:
                upload_ok = self.storage.upload_fileobj(
                    f,
                    docx_key,
                    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            if not upload_ok:
                raise RuntimeError(f"Failed to upload DOCX to storage: {docx_key}")

            docx_url = self.storage.generate_presigned_url(docx_key, expiration=300)

            # Attempt PDF conversion
            pdf_url: Optional[str] = None
            pdf_key: Optional[str] = None
            warning: Optional[str] = None

            pdf_path = self.convert_to_pdf(docx_path, tmpdir)
            if pdf_path:
                pdf_filename = os.path.basename(pdf_path)
                pdf_key = f"{law_firm_id}/exports/{version_id}/{pdf_filename}"
                with open(pdf_path, "rb") as f:
                    upload_ok = self.storage.upload_fileobj(
                        f,
                        pdf_key,
                        content_type="application/pdf",
                    )
                if upload_ok:
                    pdf_url = self.storage.generate_presigned_url(pdf_key, expiration=300)
                else:
                    warning = "PDF conversion succeeded but upload to storage failed. DOCX export is available."
            else:
                warning = (
                    "PDF conversion unavailable: LibreOffice is not installed or conversion failed. "
                    "Only DOCX export is available."
                )

        return {
            "docx_url": docx_url,
            "pdf_url": pdf_url,
            "docx_storage_key": docx_key,
            "pdf_storage_key": pdf_key,
            "warning": warning,
        }
